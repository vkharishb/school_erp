from __future__ import annotations

from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Annotated, Any
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import and_, case, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import (
    ensure_campus_access,
    ensure_license_valid,
    ensure_school_access,
    require_permissions,
)
from app.db.session import get_db
from app.models.academic import AcademicClass, Section, Student, StudentEnrollment, Subject
from app.models.attendance import StudentAttendance, TeacherAttendance
from app.models.fee import FeeHead, Payment, PaymentAllocation, StudentCharge
from app.models.marks import StudentMark
from app.models.organization import AcademicYear, Campus
from app.models.staff import Teacher
from app.models.user import User
from app.schemas.reports import ReportColumn, ReportResult
from app.services.audit import record_audit

router = APIRouter(prefix="/reports", tags=["Reports"])

REPORTS: dict[str, dict[str, str]] = {
    "student_roster": {
        "title": "Student Enrollment Report",
        "permission": "student.view",
        "module": "student",
    },
    "teacher_roster": {
        "title": "Teacher Report",
        "permission": "teacher.view",
        "module": "teacher",
    },
    "student_attendance": {
        "title": "Student Attendance Summary",
        "permission": "attendance.view",
        "module": "attendance",
    },
    "teacher_attendance": {
        "title": "Teacher Attendance Summary",
        "permission": "attendance.view",
        "module": "attendance",
    },
    "fee_collections": {
        "title": "Fee Collection Report",
        "permission": "fee.view",
        "module": "fee",
    },
    "fee_dues": {"title": "Outstanding Fee Dues Report", "permission": "fee.view", "module": "fee"},
    "marks_performance": {
        "title": "Marks Performance Report",
        "permission": "marks.view",
        "module": "marks",
    },
}

VIEW_LIMIT_MAX = 1000
EXPORT_LIMIT_MAX = 10000
MAX_DATE_RANGE_DAYS = 366


def _permission_codes(user: User) -> set[str]:
    if user.is_superuser:
        return {"*"}
    return {
        rp.permission.code
        for assignment in user.roles
        if assignment.role
        for rp in assignment.role.permissions
        if rp.permission
    }


def _require_permission(user: User, code: str) -> None:
    if user.is_superuser:
        return
    permissions = _permission_codes(user)
    if "*" not in permissions and code not in permissions:
        raise HTTPException(status_code=403, detail=f"Missing permission: {code}")


def _validate_report_key(report_key: str) -> dict[str, str]:
    definition = REPORTS.get(report_key)
    if not definition:
        raise HTTPException(status_code=404, detail="Unknown report")
    return definition


def _date_window(start_date: date | None, end_date: date | None) -> tuple[date, date]:
    today = datetime.now(UTC).date()
    end = end_date or today
    start = start_date or (end - timedelta(days=29))
    if start > end:
        raise HTTPException(status_code=422, detail="Start date cannot be after end date")
    if (end - start).days > MAX_DATE_RANGE_DAYS:
        raise HTTPException(
            status_code=422, detail=f"Report date range cannot exceed {MAX_DATE_RANGE_DAYS} days"
        )
    return start, end


def _datetime_window(start_date: date | None, end_date: date | None) -> tuple[datetime, datetime]:
    start, end = _date_window(start_date, end_date)
    return (
        datetime.combine(start, time.min, tzinfo=UTC),
        datetime.combine(end + timedelta(days=1), time.min, tzinfo=UTC),
    )


def _name(first: str | None, last: str | None) -> str:
    return " ".join(part for part in [first or "", last or ""] if part).strip()


def _money(value: Any) -> str:
    return f"{Decimal(str(value or 0)):.2f}"


def _safe_cell(value: Any) -> Any:
    # Prevent spreadsheet formula injection in exported user-controlled text.
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


async def _authorize(
    db: AsyncSession,
    user: User,
    school_id: UUID,
    report_key: str,
    campus_id: UUID | None,
    *,
    export: bool = False,
) -> UUID | None:
    definition = _validate_report_key(report_key)
    await ensure_school_access(user, db, school_id)
    await ensure_license_valid(school_id, db, "reports")
    await ensure_license_valid(school_id, db, definition["module"])
    _require_permission(user, definition["permission"])
    if export:
        _require_permission(user, "reports.export")
        if report_key.startswith("fee_"):
            _require_permission(user, "fee.export")

    if user.campus_id:
        if campus_id and campus_id != user.campus_id:
            raise HTTPException(status_code=403, detail="Campus access denied")
        campus_id = user.campus_id
    if campus_id:
        campus = await ensure_campus_access(user, db, campus_id)
        if campus.school_id != school_id:
            raise HTTPException(status_code=403, detail="Campus does not belong to this school")
    return campus_id


async def _validate_filter_scope(
    db: AsyncSession,
    school_id: UUID,
    campus_id: UUID | None,
    academic_year_id: UUID | None,
    academic_class_id: UUID | None,
    section_id: UUID | None,
    subject_id: UUID | None,
) -> None:
    if academic_year_id:
        year = await db.get(AcademicYear, academic_year_id)
        if not year:
            raise HTTPException(status_code=404, detail="Academic year not found")
        campus = await db.get(Campus, year.campus_id)
        if not campus or campus.school_id != school_id or (campus_id and campus.id != campus_id):
            raise HTTPException(
                status_code=422, detail="Academic year is outside the selected school/campus"
            )
    if academic_class_id:
        academic_class = await db.get(AcademicClass, academic_class_id)
        if not academic_class:
            raise HTTPException(status_code=404, detail="Class not found")
        campus = await db.get(Campus, academic_class.campus_id)
        if not campus or campus.school_id != school_id or (campus_id and campus.id != campus_id):
            raise HTTPException(
                status_code=422, detail="Class is outside the selected school/campus"
            )
    if section_id:
        section = await db.get(Section, section_id)
        if not section:
            raise HTTPException(status_code=404, detail="Section not found")
        academic_class = await db.get(AcademicClass, section.academic_class_id)
        if not academic_class or (academic_class_id and academic_class.id != academic_class_id):
            raise HTTPException(
                status_code=422, detail="Section does not belong to the selected class"
            )
        campus = await db.get(Campus, academic_class.campus_id)
        if not campus or campus.school_id != school_id or (campus_id and campus.id != campus_id):
            raise HTTPException(
                status_code=422, detail="Section is outside the selected school/campus"
            )
    if subject_id:
        subject = await db.get(Subject, subject_id)
        if not subject:
            raise HTTPException(status_code=404, detail="Subject not found")
        campus = await db.get(Campus, subject.campus_id)
        if not campus or campus.school_id != school_id or (campus_id and campus.id != campus_id):
            raise HTTPException(
                status_code=422, detail="Subject is outside the selected school/campus"
            )


def _result(
    report_key: str,
    columns: list[tuple[str, str]],
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    total_count: int,
    limit: int,
) -> ReportResult:
    return ReportResult(
        report_key=report_key,
        title=REPORTS[report_key]["title"],
        columns=[ReportColumn(key=key, label=label) for key, label in columns],
        rows=rows,
        summary=summary,
        total_count=total_count,
        generated_at=datetime.now(UTC),
        truncated=total_count > limit,
    )


async def _student_roster(
    db: AsyncSession,
    school_id: UUID,
    campus_id: UUID | None,
    academic_year_id: UUID | None,
    academic_class_id: UUID | None,
    section_id: UUID | None,
    status: str | None,
    limit: int,
) -> ReportResult:
    conditions = [Student.school_id == school_id]
    if campus_id:
        conditions.append(Student.campus_id == campus_id)
    if academic_year_id:
        conditions.append(StudentEnrollment.academic_year_id == academic_year_id)
    if academic_class_id:
        conditions.append(StudentEnrollment.academic_class_id == academic_class_id)
    if section_id:
        conditions.append(StudentEnrollment.section_id == section_id)
    if status:
        conditions.append(Student.status == status)

    base = (
        select(
            Student.admission_number,
            Student.first_name,
            Student.last_name,
            Student.status.label("student_status"),
            Campus.name.label("campus_name"),
            AcademicYear.name.label("year_name"),
            AcademicClass.name.label("class_name"),
            Section.name.label("section_name"),
            StudentEnrollment.roll_number,
            StudentEnrollment.status.label("enrollment_status"),
        )
        .join(StudentEnrollment, StudentEnrollment.student_id == Student.id)
        .join(AcademicYear, AcademicYear.id == StudentEnrollment.academic_year_id)
        .join(AcademicClass, AcademicClass.id == StudentEnrollment.academic_class_id)
        .join(Section, Section.id == StudentEnrollment.section_id)
        .join(Campus, Campus.id == Student.campus_id)
        .where(*conditions)
    )
    total = int((await db.execute(select(func.count()).select_from(base.subquery()))).scalar_one())
    result = await db.execute(
        base.order_by(
            Campus.name,
            AcademicClass.sort_order,
            Section.name,
            StudentEnrollment.roll_number.nulls_last(),
            Student.first_name,
        ).limit(limit)
    )
    rows = [
        {
            "admission_number": r.admission_number,
            "student_name": _name(r.first_name, r.last_name),
            "campus": r.campus_name,
            "academic_year": r.year_name,
            "class": r.class_name,
            "section": r.section_name,
            "roll_number": r.roll_number or "",
            "student_status": r.student_status,
            "enrollment_status": r.enrollment_status,
        }
        for r in result.all()
    ]
    return _result(
        "student_roster",
        [
            ("admission_number", "Admission No."),
            ("student_name", "Student Name"),
            ("campus", "Campus"),
            ("academic_year", "Academic Year"),
            ("class", "Class"),
            ("section", "Section"),
            ("roll_number", "Roll No."),
            ("student_status", "Student Status"),
            ("enrollment_status", "Enrollment Status"),
        ],
        rows,
        {"matching_students": total},
        total,
        limit,
    )


async def _teacher_roster(
    db: AsyncSession, school_id: UUID, campus_id: UUID | None, status: str | None, limit: int
) -> ReportResult:
    conditions = [Teacher.school_id == school_id]
    if campus_id:
        conditions.append(Teacher.campus_id == campus_id)
    if status:
        conditions.append(Teacher.status == status)
    base = (
        select(
            Teacher.employee_code,
            Teacher.first_name,
            Teacher.last_name,
            Teacher.designation,
            Teacher.qualification,
            Teacher.category,
            Teacher.joining_date,
            Teacher.status,
            Campus.name.label("campus_name"),
        )
        .join(Campus, Campus.id == Teacher.campus_id)
        .where(*conditions)
    )
    total = int((await db.execute(select(func.count()).select_from(base.subquery()))).scalar_one())
    result = await db.execute(
        base.order_by(Campus.name, Teacher.first_name, Teacher.employee_code).limit(limit)
    )
    rows = [
        {
            "employee_code": r.employee_code,
            "teacher_name": _name(r.first_name, r.last_name),
            "campus": r.campus_name,
            "designation": r.designation or "",
            "qualification": r.qualification or "",
            "category": r.category or "",
            "joining_date": r.joining_date.isoformat() if r.joining_date else "",
            "status": r.status,
        }
        for r in result.all()
    ]
    return _result(
        "teacher_roster",
        [
            ("employee_code", "Employee No."),
            ("teacher_name", "Teacher Name"),
            ("campus", "Campus"),
            ("designation", "Designation"),
            ("qualification", "Qualification"),
            ("category", "Category"),
            ("joining_date", "Joining Date"),
            ("status", "Status"),
        ],
        rows,
        {"matching_teachers": total},
        total,
        limit,
    )


async def _student_attendance(
    db: AsyncSession,
    school_id: UUID,
    campus_id: UUID | None,
    academic_year_id: UUID | None,
    academic_class_id: UUID | None,
    section_id: UUID | None,
    start_date: date | None,
    end_date: date | None,
    limit: int,
) -> ReportResult:
    start, end = _date_window(start_date, end_date)
    conditions = [
        StudentAttendance.school_id == school_id,
        StudentAttendance.attendance_date >= start,
        StudentAttendance.attendance_date <= end,
    ]
    if campus_id:
        conditions.append(StudentAttendance.campus_id == campus_id)
    include_enrollment = bool(academic_year_id or academic_class_id or section_id)
    if (academic_class_id or section_id) and not academic_year_id:
        raise HTTPException(
            status_code=422,
            detail="Select an Academic Year before filtering attendance by Class/Section",
        )
    columns = [
        ("admission_number", "Admission No."),
        ("student_name", "Student Name"),
        ("campus", "Campus"),
    ]
    group_fields: list[Any] = [
        Student.admission_number,
        Student.first_name,
        Student.last_name,
        Campus.name,
    ]
    select_fields: list[Any] = [
        Student.admission_number,
        Student.first_name,
        Student.last_name,
        Campus.name.label("campus_name"),
    ]
    query = select(*select_fields)
    if include_enrollment:
        query = query.add_columns(
            AcademicClass.name.label("class_name"), Section.name.label("section_name")
        )
        group_fields += [AcademicClass.name, Section.name]
        columns += [("class", "Class"), ("section", "Section")]
    present = func.sum(case((StudentAttendance.status == "present", 1), else_=0)).label("present")
    absent = func.sum(case((StudentAttendance.status == "absent", 1), else_=0)).label("absent")
    late = func.sum(case((StudentAttendance.status == "late", 1), else_=0)).label("late")
    half_day = func.sum(case((StudentAttendance.status == "half_day", 1), else_=0)).label(
        "half_day"
    )
    leave = func.sum(case((StudentAttendance.status == "leave", 1), else_=0)).label("leave")
    total_days = func.count(StudentAttendance.id).label("days_marked")
    query = (
        query.add_columns(total_days, present, absent, late, half_day, leave)
        .select_from(StudentAttendance)
        .join(Student, Student.id == StudentAttendance.student_id)
        .join(Campus, Campus.id == StudentAttendance.campus_id)
    )
    if include_enrollment:
        query = (
            query.join(
                StudentEnrollment,
                and_(
                    StudentEnrollment.student_id == Student.id,
                    StudentEnrollment.academic_year_id == academic_year_id,
                ),
            )
            .join(AcademicClass, AcademicClass.id == StudentEnrollment.academic_class_id)
            .join(Section, Section.id == StudentEnrollment.section_id)
        )
        if academic_class_id:
            conditions.append(StudentEnrollment.academic_class_id == academic_class_id)
        if section_id:
            conditions.append(StudentEnrollment.section_id == section_id)
    query = query.where(*conditions).group_by(*group_fields)
    total = int((await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one())
    data = (await db.execute(query.order_by(Campus.name, Student.first_name).limit(limit))).all()
    rows: list[dict[str, Any]] = []
    totals = {"present": 0, "absent": 0, "late": 0, "half_day": 0, "leave": 0, "days_marked": 0}
    for r in data:
        days = int(r.days_marked or 0)
        p = int(r.present or 0)
        a = int(r.absent or 0)
        late = int(r.late or 0)
        h = int(r.half_day or 0)
        lv = int(r.leave or 0)
        row = {
            "admission_number": r.admission_number,
            "student_name": _name(r.first_name, r.last_name),
            "campus": r.campus_name,
        }
        if include_enrollment:
            row.update({"class": r.class_name, "section": r.section_name})
        row.update(
            {
                "days_marked": days,
                "present": p,
                "absent": a,
                "late": late,
                "half_day": h,
                "leave": lv,
                "attendance_pct": f"{((p + (h * 0.5)) / days * 100) if days else 0:.2f}",
            }
        )
        rows.append(row)
        totals["days_marked"] += days
        totals["present"] += p
        totals["absent"] += a
        totals["late"] += late
        totals["half_day"] += h
        totals["leave"] += lv
    columns += [
        ("days_marked", "Days Marked"),
        ("present", "Present"),
        ("absent", "Absent"),
        ("late", "Late"),
        ("half_day", "Half Day"),
        ("leave", "Leave"),
        ("attendance_pct", "Attendance %"),
    ]
    return _result(
        "student_attendance",
        columns,
        rows,
        {"students": total, "date_from": start.isoformat(), "date_to": end.isoformat(), **totals},
        total,
        limit,
    )


async def _teacher_attendance(
    db: AsyncSession,
    school_id: UUID,
    campus_id: UUID | None,
    start_date: date | None,
    end_date: date | None,
    limit: int,
) -> ReportResult:
    start, end = _date_window(start_date, end_date)
    conditions = [
        TeacherAttendance.school_id == school_id,
        TeacherAttendance.attendance_date >= start,
        TeacherAttendance.attendance_date <= end,
    ]
    if campus_id:
        conditions.append(TeacherAttendance.campus_id == campus_id)
    present = func.sum(case((TeacherAttendance.status == "present", 1), else_=0)).label("present")
    absent = func.sum(case((TeacherAttendance.status == "absent", 1), else_=0)).label("absent")
    late = func.sum(case((TeacherAttendance.status == "late", 1), else_=0)).label("late")
    half_day = func.sum(case((TeacherAttendance.status == "half_day", 1), else_=0)).label(
        "half_day"
    )
    leave = func.sum(case((TeacherAttendance.status == "leave", 1), else_=0)).label("leave")
    days = func.count(TeacherAttendance.id).label("days_marked")
    query = (
        select(
            Teacher.employee_code,
            Teacher.first_name,
            Teacher.last_name,
            Campus.name.label("campus_name"),
            days,
            present,
            absent,
            late,
            half_day,
            leave,
        )
        .select_from(TeacherAttendance)
        .join(Teacher, Teacher.id == TeacherAttendance.teacher_id)
        .join(Campus, Campus.id == TeacherAttendance.campus_id)
        .where(*conditions)
        .group_by(Teacher.employee_code, Teacher.first_name, Teacher.last_name, Campus.name)
    )
    total = int((await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one())
    data = (await db.execute(query.order_by(Campus.name, Teacher.first_name).limit(limit))).all()
    rows = []
    totals = {"present": 0, "absent": 0, "late": 0, "half_day": 0, "leave": 0, "days_marked": 0}
    for r in data:
        d = int(r.days_marked or 0)
        p = int(r.present or 0)
        a = int(r.absent or 0)
        late = int(r.late or 0)
        h = int(r.half_day or 0)
        lv = int(r.leave or 0)
        rows.append(
            {
                "employee_code": r.employee_code,
                "teacher_name": _name(r.first_name, r.last_name),
                "campus": r.campus_name,
                "days_marked": d,
                "present": p,
                "absent": a,
                "late": late,
                "half_day": h,
                "leave": lv,
                "attendance_pct": f"{((p + (h * 0.5)) / d * 100) if d else 0:.2f}",
            }
        )
        totals["days_marked"] += d
        totals["present"] += p
        totals["absent"] += a
        totals["late"] += late
        totals["half_day"] += h
        totals["leave"] += lv
    return _result(
        "teacher_attendance",
        [
            ("employee_code", "Employee No."),
            ("teacher_name", "Teacher Name"),
            ("campus", "Campus"),
            ("days_marked", "Days Marked"),
            ("present", "Present"),
            ("absent", "Absent"),
            ("late", "Late"),
            ("half_day", "Half Day"),
            ("leave", "Leave"),
            ("attendance_pct", "Attendance %"),
        ],
        rows,
        {"teachers": total, "date_from": start.isoformat(), "date_to": end.isoformat(), **totals},
        total,
        limit,
    )


async def _fee_collections(
    db: AsyncSession,
    school_id: UUID,
    campus_id: UUID | None,
    start_date: date | None,
    end_date: date | None,
    payment_mode: str | None,
    status: str | None,
    limit: int,
) -> ReportResult:
    start_dt, end_dt = _datetime_window(start_date, end_date)
    conditions = [
        Payment.school_id == school_id,
        Payment.collected_at >= start_dt,
        Payment.collected_at < end_dt,
    ]
    if campus_id:
        conditions.append(Payment.campus_id == campus_id)
    if payment_mode:
        conditions.append(func.lower(Payment.payment_mode) == payment_mode.lower())
    if status:
        conditions.append(Payment.status == status)
    base = (
        select(
            Payment.receipt_number,
            Payment.collected_at,
            Payment.amount,
            Payment.payment_mode,
            Payment.status,
            Student.admission_number,
            Student.first_name,
            Student.last_name,
            Campus.name.label("campus_name"),
        )
        .join(Student, Student.id == Payment.student_id)
        .join(Campus, Campus.id == Payment.campus_id)
        .where(*conditions)
    )
    total = int((await db.execute(select(func.count()).select_from(base.subquery()))).scalar_one())
    agg = await db.execute(
        select(
            func.coalesce(
                func.sum(case((Payment.status == "posted", Payment.amount), else_=0)), 0
            ).label("posted_total"),
            func.coalesce(func.sum(case((Payment.status != "posted", 1), else_=0)), 0).label(
                "non_posted"
            ),
        ).where(*conditions)
    )
    agg_row = agg.one()
    data = (await db.execute(base.order_by(Payment.collected_at.desc()).limit(limit))).all()
    rows = [
        {
            "receipt_number": r.receipt_number,
            "collected_at": r.collected_at.isoformat(),
            "admission_number": r.admission_number,
            "student_name": _name(r.first_name, r.last_name),
            "campus": r.campus_name,
            "payment_mode": r.payment_mode,
            "amount": _money(r.amount),
            "status": r.status,
        }
        for r in data
    ]
    return _result(
        "fee_collections",
        [
            ("receipt_number", "Receipt No."),
            ("collected_at", "Collected At"),
            ("admission_number", "Admission No."),
            ("student_name", "Student Name"),
            ("campus", "Campus"),
            ("payment_mode", "Payment Mode"),
            ("amount", "Amount"),
            ("status", "Status"),
        ],
        rows,
        {
            "receipts": total,
            "posted_amount": _money(agg_row.posted_total),
            "non_posted_receipts": int(agg_row.non_posted or 0),
        },
        total,
        limit,
    )


async def _fee_dues(
    db: AsyncSession,
    school_id: UUID,
    campus_id: UUID | None,
    academic_year_id: UUID | None,
    academic_class_id: UUID | None,
    section_id: UUID | None,
    limit: int,
) -> ReportResult:
    paid_sq = (
        select(
            PaymentAllocation.charge_id.label("charge_id"),
            func.coalesce(func.sum(PaymentAllocation.amount), 0).label("paid"),
        )
        .join(Payment, Payment.id == PaymentAllocation.payment_id)
        .where(Payment.status == "posted")
        .group_by(PaymentAllocation.charge_id)
        .subquery()
    )
    paid = func.coalesce(paid_sq.c.paid, 0)
    outstanding = (StudentCharge.amount - paid).label("outstanding")
    conditions = [StudentCharge.school_id == school_id, StudentCharge.amount - paid > 0]
    if campus_id:
        conditions.append(Student.campus_id == campus_id)
    if academic_year_id:
        conditions.append(StudentCharge.academic_year_id == academic_year_id)
    include_enrollment = bool(academic_class_id or section_id)
    if include_enrollment and not academic_year_id:
        raise HTTPException(
            status_code=422, detail="Select an Academic Year before filtering dues by Class/Section"
        )
    query = (
        select(
            StudentCharge.id,
            Student.admission_number,
            Student.first_name,
            Student.last_name,
            Campus.name.label("campus_name"),
            FeeHead.name.label("fee_head"),
            StudentCharge.description,
            StudentCharge.amount,
            paid.label("paid"),
            outstanding,
            StudentCharge.due_date,
            StudentCharge.status,
        )
        .select_from(StudentCharge)
        .join(Student, Student.id == StudentCharge.student_id)
        .join(Campus, Campus.id == Student.campus_id)
        .join(FeeHead, FeeHead.id == StudentCharge.fee_head_id)
        .outerjoin(paid_sq, paid_sq.c.charge_id == StudentCharge.id)
    )
    if include_enrollment:
        query = (
            query.join(
                StudentEnrollment,
                and_(
                    StudentEnrollment.student_id == Student.id,
                    StudentEnrollment.academic_year_id == academic_year_id,
                ),
            )
            .join(AcademicClass, AcademicClass.id == StudentEnrollment.academic_class_id)
            .join(Section, Section.id == StudentEnrollment.section_id)
            .add_columns(AcademicClass.name.label("class_name"), Section.name.label("section_name"))
        )
        if academic_class_id:
            conditions.append(StudentEnrollment.academic_class_id == academic_class_id)
        if section_id:
            conditions.append(StudentEnrollment.section_id == section_id)
    query = query.where(*conditions)
    total = int((await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one())
    data = (
        await db.execute(
            query.order_by(StudentCharge.due_date.nulls_last(), Student.admission_number).limit(
                limit
            )
        )
    ).all()
    rows = []
    total_due = Decimal("0")
    for r in data:
        row = {
            "admission_number": r.admission_number,
            "student_name": _name(r.first_name, r.last_name),
            "campus": r.campus_name,
            "fee_head": r.fee_head,
            "description": r.description,
            "charge_amount": _money(r.amount),
            "paid": _money(r.paid),
            "outstanding": _money(r.outstanding),
            "due_date": r.due_date.isoformat() if r.due_date else "",
            "status": r.status,
        }
        if include_enrollment:
            row.update({"class": r.class_name, "section": r.section_name})
        rows.append(row)
        total_due += Decimal(str(r.outstanding or 0))
    columns = [
        ("admission_number", "Admission No."),
        ("student_name", "Student Name"),
        ("campus", "Campus"),
    ]
    if include_enrollment:
        columns += [("class", "Class"), ("section", "Section")]
    columns += [
        ("fee_head", "Fee Head"),
        ("description", "Description"),
        ("charge_amount", "Charge"),
        ("paid", "Paid"),
        ("outstanding", "Outstanding"),
        ("due_date", "Due Date"),
        ("status", "Status"),
    ]
    # Exact total across all matching rows, not just the UI page.
    total_amount = (
        await db.execute(select(func.coalesce(func.sum(query.subquery().c.outstanding), 0)))
    ).scalar_one()
    return _result(
        "fee_dues",
        columns,
        rows,
        {"outstanding_items": total, "total_outstanding": _money(total_amount)},
        total,
        limit,
    )


async def _marks_performance(
    db: AsyncSession,
    school_id: UUID,
    campus_id: UUID | None,
    academic_year_id: UUID | None,
    academic_class_id: UUID | None,
    section_id: UUID | None,
    subject_id: UUID | None,
    assessment_name: str | None,
    limit: int,
) -> ReportResult:
    conditions = [StudentMark.school_id == school_id]
    if campus_id:
        conditions.append(StudentMark.campus_id == campus_id)
    if academic_year_id:
        conditions.append(StudentMark.academic_year_id == academic_year_id)
    if academic_class_id:
        conditions.append(StudentEnrollment.academic_class_id == academic_class_id)
    if section_id:
        conditions.append(StudentEnrollment.section_id == section_id)
    if subject_id:
        conditions.append(StudentMark.subject_id == subject_id)
    if assessment_name:
        conditions.append(
            func.lower(StudentMark.assessment_name) == assessment_name.strip().lower()
        )
    query = (
        select(
            Student.admission_number,
            Student.first_name,
            Student.last_name,
            Campus.name.label("campus_name"),
            AcademicYear.name.label("year_name"),
            AcademicClass.name.label("class_name"),
            Section.name.label("section_name"),
            Subject.name.label("subject_name"),
            StudentMark.assessment_name,
            StudentMark.max_marks,
            StudentMark.marks_obtained,
            StudentMark.is_locked,
        )
        .select_from(StudentMark)
        .join(Student, Student.id == StudentMark.student_id)
        .join(Campus, Campus.id == StudentMark.campus_id)
        .join(AcademicYear, AcademicYear.id == StudentMark.academic_year_id)
        .join(Subject, Subject.id == StudentMark.subject_id)
        .join(
            StudentEnrollment,
            and_(
                StudentEnrollment.student_id == Student.id,
                StudentEnrollment.academic_year_id == StudentMark.academic_year_id,
            ),
        )
        .join(AcademicClass, AcademicClass.id == StudentEnrollment.academic_class_id)
        .join(Section, Section.id == StudentEnrollment.section_id)
        .where(*conditions)
    )
    total = int((await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one())
    data = (
        await db.execute(
            query.order_by(AcademicClass.sort_order, Section.name, Student.admission_number).limit(
                limit
            )
        )
    ).all()
    rows = []
    for r in data:
        max_marks = Decimal(str(r.max_marks or 0))
        obtained = Decimal(str(r.marks_obtained or 0))
        pct = (obtained / max_marks * Decimal("100")) if max_marks else Decimal("0")
        rows.append(
            {
                "admission_number": r.admission_number,
                "student_name": _name(r.first_name, r.last_name),
                "campus": r.campus_name,
                "academic_year": r.year_name,
                "class": r.class_name,
                "section": r.section_name,
                "subject": r.subject_name,
                "assessment": r.assessment_name,
                "max_marks": _money(max_marks),
                "marks_obtained": _money(obtained),
                "percentage": f"{pct:.2f}",
                "locked": "Yes" if r.is_locked else "No",
            }
        )
    percentage_expr = StudentMark.marks_obtained / func.nullif(StudentMark.max_marks, 0) * 100
    agg_query = (
        select(
            func.count(StudentMark.id),
            func.avg(percentage_expr),
            func.max(percentage_expr),
            func.min(percentage_expr),
        )
        .select_from(StudentMark)
        .join(Student, Student.id == StudentMark.student_id)
        .join(
            StudentEnrollment,
            and_(
                StudentEnrollment.student_id == Student.id,
                StudentEnrollment.academic_year_id == StudentMark.academic_year_id,
            ),
        )
        .where(*conditions)
    )
    agg = (await db.execute(agg_query)).one()
    return _result(
        "marks_performance",
        [
            ("admission_number", "Admission No."),
            ("student_name", "Student Name"),
            ("campus", "Campus"),
            ("academic_year", "Academic Year"),
            ("class", "Class"),
            ("section", "Section"),
            ("subject", "Subject"),
            ("assessment", "Assessment"),
            ("max_marks", "Max Marks"),
            ("marks_obtained", "Marks"),
            ("percentage", "Percentage"),
            ("locked", "Finalized"),
        ],
        rows,
        {
            "marks_records": int(agg[0] or 0),
            "average_percentage": f"{Decimal(str(agg[1] or 0)):.2f}",
            "highest_percentage": f"{Decimal(str(agg[2] or 0)):.2f}",
            "lowest_percentage": f"{Decimal(str(agg[3] or 0)):.2f}",
        },
        total,
        limit,
    )


async def _build_report(
    db: AsyncSession,
    school_id: UUID,
    report_key: str,
    *,
    campus_id: UUID | None,
    academic_year_id: UUID | None,
    academic_class_id: UUID | None,
    section_id: UUID | None,
    subject_id: UUID | None,
    assessment_name: str | None,
    start_date: date | None,
    end_date: date | None,
    status: str | None,
    payment_mode: str | None,
    limit: int,
) -> ReportResult:
    await _validate_filter_scope(
        db, school_id, campus_id, academic_year_id, academic_class_id, section_id, subject_id
    )
    if report_key == "student_roster":
        return await _student_roster(
            db, school_id, campus_id, academic_year_id, academic_class_id, section_id, status, limit
        )
    if report_key == "teacher_roster":
        return await _teacher_roster(db, school_id, campus_id, status, limit)
    if report_key == "student_attendance":
        return await _student_attendance(
            db,
            school_id,
            campus_id,
            academic_year_id,
            academic_class_id,
            section_id,
            start_date,
            end_date,
            limit,
        )
    if report_key == "teacher_attendance":
        return await _teacher_attendance(db, school_id, campus_id, start_date, end_date, limit)
    if report_key == "fee_collections":
        return await _fee_collections(
            db, school_id, campus_id, start_date, end_date, payment_mode, status, limit
        )
    if report_key == "fee_dues":
        return await _fee_dues(
            db, school_id, campus_id, academic_year_id, academic_class_id, section_id, limit
        )
    if report_key == "marks_performance":
        return await _marks_performance(
            db,
            school_id,
            campus_id,
            academic_year_id,
            academic_class_id,
            section_id,
            subject_id,
            assessment_name,
            limit,
        )
    raise HTTPException(status_code=404, detail="Unknown report")


@router.get("/{school_id}/run/{report_key}", response_model=ReportResult)
async def run_report(
    school_id: UUID,
    report_key: str,
    db: Annotated[AsyncSession, Depends(get_db)],
    user: Annotated[User, Depends(require_permissions("reports.view"))],
    campus_id: UUID | None = None,
    academic_year_id: UUID | None = None,
    academic_class_id: UUID | None = None,
    section_id: UUID | None = None,
    subject_id: UUID | None = None,
    assessment_name: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    status: str | None = None,
    payment_mode: str | None = None,
    limit: int = Query(500, ge=1, le=VIEW_LIMIT_MAX),
):
    campus_id = await _authorize(db, user, school_id, report_key, campus_id)
    return await _build_report(
        db,
        school_id,
        report_key,
        campus_id=campus_id,
        academic_year_id=academic_year_id,
        academic_class_id=academic_class_id,
        section_id=section_id,
        subject_id=subject_id,
        assessment_name=assessment_name,
        start_date=start_date,
        end_date=end_date,
        status=status,
        payment_mode=payment_mode,
        limit=limit,
    )


@router.get("/{school_id}/export/{report_key}")
async def export_report(
    school_id: UUID,
    report_key: str,
    request: Request,
    db: Annotated[AsyncSession, Depends(get_db)],
    user: Annotated[User, Depends(require_permissions("reports.view", "reports.export"))],
    campus_id: UUID | None = None,
    academic_year_id: UUID | None = None,
    academic_class_id: UUID | None = None,
    section_id: UUID | None = None,
    subject_id: UUID | None = None,
    assessment_name: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    status: str | None = None,
    payment_mode: str | None = None,
):
    campus_id = await _authorize(db, user, school_id, report_key, campus_id, export=True)
    report = await _build_report(
        db,
        school_id,
        report_key,
        campus_id=campus_id,
        academic_year_id=academic_year_id,
        academic_class_id=academic_class_id,
        section_id=section_id,
        subject_id=subject_id,
        assessment_name=assessment_name,
        start_date=start_date,
        end_date=end_date,
        status=status,
        payment_mode=payment_mode,
        limit=EXPORT_LIMIT_MAX,
    )
    if report.total_count > EXPORT_LIMIT_MAX:
        raise HTTPException(
            status_code=413,
            detail=f"Report has {report.total_count} rows. Narrow filters before export; maximum is {EXPORT_LIMIT_MAX} rows.",
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Generated", report.generated_at.isoformat()])
    for key, value in report.summary.items():
        ws.append([str(key).replace("_", " ").title(), _safe_cell(value)])
    ws.append([])
    ws.append([column.label for column in report.columns])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for row in report.rows:
        ws.append([_safe_cell(row.get(column.key, "")) for column in report.columns])
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = (
        f"A{header_row}:{chr(64 + min(len(report.columns), 26))}{ws.max_row}"
        if len(report.columns) <= 26
        else None
    )
    for column_cells in ws.columns:
        max_len = min(
            max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2, 40
        )
        ws.column_dimensions[column_cells[0].column_letter].width = max_len

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    await record_audit(
        db,
        action="report.exported",
        user=user,
        module="reports",
        entity_type="Report",
        after={
            "report_key": report_key,
            "school_id": str(school_id),
            "row_count": report.total_count,
            "campus_id": str(campus_id) if campus_id else None,
        },
        request=request,
    )
    filename = f"{report_key}-{datetime.now(UTC).date().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )
