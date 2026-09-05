from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_ROWS = 2000
MAX_PREVIEW_ROWS = 20


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", text)).strip("_")


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def parse_date(value: Any, *, field: str) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"{field} must be a valid Excel date or YYYY-MM-DD") from exc


def parse_decimal(value: Any, *, field: str, required: bool = False) -> Decimal | None:
    if value is None or value == "":
        if required:
            raise ValueError(f"{field} is required")
        return None
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} must be a valid number") from exc


def split_name(value: Any, *, field: str) -> tuple[str, str | None]:
    name = clean_text(value)
    if not name:
        raise ValueError(f"{field} is required")
    parts = name.split()
    return parts[0], " ".join(parts[1:]) or None


def normalize_phone(value: Any, *, field: str, required: bool = False) -> str | None:
    text = clean_text(value)
    if not text:
        if required:
            raise ValueError(f"{field} is required")
        return None
    normalized = re.sub(r"[\s()-]", "", text)
    if normalized.startswith("+"):
        digits = normalized[1:]
    else:
        digits = normalized
    if not digits.isdigit() or not 7 <= len(digits) <= 15:
        raise ValueError(f"{field} must contain 7 to 15 digits")
    return normalized


def normalize_aadhaar(value: Any, *, field: str = "Aadhaar Number") -> str | None:
    text = clean_text(value)
    if not text:
        return None
    digits = re.sub(r"\s", "", text)
    if not digits.isdigit() or len(digits) != 12:
        raise ValueError(f"{field} must contain exactly 12 digits")
    return digits


def safe_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@dataclass(frozen=True)
class ParsedWorkbook:
    sha256: str
    rows: list[dict[str, Any]]


async def read_upload(upload: UploadFile) -> bytes:
    filename = Path(upload.filename or "").name
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="Only .xlsx files are accepted")
    data = await upload.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Upload exceeds the 5 MB limit")
    if not data:
        raise HTTPException(status_code=422, detail="Uploaded workbook is empty")
    if not zipfile.is_zipfile(BytesIO(data)):
        raise HTTPException(status_code=422, detail="Uploaded file is not a valid XLSX workbook")
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > 1000:
                raise HTTPException(
                    status_code=422, detail="Workbook contains too many embedded parts"
                )
            if sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_BYTES:
                raise HTTPException(
                    status_code=422, detail="Workbook expands beyond the safe processing limit"
                )
            names = {info.filename.replace("\\", "/").lower() for info in infos}
            if any(
                name.endswith("vbaproject.bin")
                or name.startswith("xl/embeddings/")
                or name.startswith("xl/externallinks/")
                for name in names
            ):
                raise HTTPException(
                    status_code=422,
                    detail="Workbooks with macros, embedded objects, or external links are not accepted",
                )
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=422, detail="Uploaded workbook is corrupt") from exc
    return data


def parse_xlsx(
    data: bytes,
    *,
    sheet_name: str,
    required_headers: set[str],
    header_aliases: dict[str, set[str]] | None = None,
) -> ParsedWorkbook:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Unable to read XLSX workbook") from exc

    if sheet_name not in workbook.sheetnames:
        raise HTTPException(
            status_code=422, detail=f"Workbook must contain a sheet named '{sheet_name}'"
        )
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows()
    try:
        header_cells = next(iterator)
    except StopIteration as exc:
        raise HTTPException(status_code=422, detail="Workbook has no header row") from exc

    headers = [normalize_header(cell.value) for cell in header_cells]
    if header_aliases:
        alias_to_canonical = {
            normalize_header(alias): canonical
            for canonical, aliases in header_aliases.items()
            for alias in ({canonical} | set(aliases))
        }
        headers = [alias_to_canonical.get(header, header) for header in headers]
    if len(headers) != len(set(h for h in headers if h)):
        raise HTTPException(status_code=422, detail="Workbook contains duplicate column headers")
    missing = sorted(required_headers - set(headers))
    if missing:
        display = ", ".join(h.replace("_", " ").title() for h in missing)
        raise HTTPException(
            status_code=422, detail=f"Workbook is missing required columns: {display}"
        )

    rows: list[dict[str, Any]] = []
    for excel_row, cells in enumerate(iterator, start=2):
        if excel_row - 1 > MAX_ROWS:
            raise HTTPException(
                status_code=422, detail=f"Workbook exceeds the {MAX_ROWS} row import limit"
            )
        if any(cell.data_type == "f" for cell in cells):
            raise HTTPException(
                status_code=422,
                detail=f"Formula cells are not allowed in import data (row {excel_row})",
            )
        values = [cell.value for cell in cells]
        if not any(value not in (None, "") for value in values):
            continue
        row = {
            headers[i]: values[i] if i < len(values) else None
            for i in range(len(headers))
            if headers[i]
        }
        row["__row__"] = excel_row
        rows.append(row)
    workbook.close()
    if not rows:
        raise HTTPException(status_code=422, detail="Workbook contains no data rows")
    return ParsedWorkbook(sha256=safe_sha256(data), rows=rows)
